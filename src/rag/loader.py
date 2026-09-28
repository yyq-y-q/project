"""
摄取层（Ingestion）—— 按后缀分流读入

知识库支持（Config.RAG_DOC_SUFFIXES）：
    .txt / .md / .markdown / .docx / .pdf（文字层）

分流约定：
    非结构化叙述 → RAG（本 loader）
    .xlsx / .csv  → sqlite_tool.ingest
    扫描版 PDF    → 无文本层则跳过（OCR 更后）

数据契约：
    documents[i]     : str
    doc_metadata[i]  : dict
        doc_id, source, source_path, file_type
"""
from __future__ import annotations

import hashlib
import logging
from pathlib import Path

from .config import Config

logger = logging.getLogger(__name__)

_SUPPORTED = {s.lower() for s in Config.RAG_DOC_SUFFIXES}
_TABLE_SUFFIXES = {".xlsx", ".xlsm", ".csv"}
_TEXT_SUFFIXES = {".txt", ".md", ".markdown"}


def _stable_doc_id(path: Path) -> str:
    abs_str = str(path.resolve())
    return hashlib.sha256(abs_str.encode("utf-8")).hexdigest()[:12]


def _read_text_file(path: Path) -> str | None:
    try:
        text = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        try:
            text = path.read_text(encoding="gbk")
        except Exception as e:
            logger.warning("读取失败，已跳过 %s: %s", path, e)
            return None
    except Exception as e:
        logger.warning("读取失败，已跳过 %s: %s", path, e)
        return None

    text = text.strip()
    if not text:
        logger.warning("空文件，已跳过: %s", path)
        return None
    return text


def _read_docx(path: Path) -> str | None:
    try:
        from docx import Document  # python-docx
    except ImportError:
        logger.error("缺少 python-docx，无法读取 %s", path)
        return None
    try:
        doc = Document(str(path))
        parts = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
        # 简单表格单元格
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = "\n".join(parts).strip()
        if not text:
            logger.warning("docx 无正文，已跳过: %s", path)
            return None
        return text
    except Exception as e:
        logger.warning("docx 解析失败，已跳过 %s: %s", path, e)
        return None


def _read_pdf(path: Path) -> str | None:
    """文字型 PDF；无文本层则跳过（不在此做 OCR）。"""
    try:
        from pypdf import PdfReader
    except ImportError:
        logger.error("缺少 pypdf，无法读取 %s", path)
        return None
    try:
        reader = PdfReader(str(path))
        if getattr(reader, "is_encrypted", False):
            try:
                reader.decrypt("")
            except Exception:
                logger.warning("加密 PDF，已跳过: %s", path)
                return None
        parts: list[str] = []
        for page in reader.pages:
            t = page.extract_text() or ""
            t = t.strip()
            if t:
                parts.append(t)
        text = "\n".join(parts).strip()
        if not text:
            logger.warning("PDF 无文字层（可能是扫描件），已跳过: %s", path)
            return None
        return text
    except Exception as e:
        logger.warning("PDF 解析失败，已跳过 %s: %s", path, e)
        return None


def _file_type(suffix: str) -> str:
    s = suffix.lower()
    if s in {".md", ".markdown"}:
        return "md"
    if s == ".docx":
        return "docx"
    if s == ".pdf":
        return "pdf"
    return "txt"


def _read_by_suffix(path: Path, suf: str) -> str | None:
    if suf in _TEXT_SUFFIXES:
        return _read_text_file(path)
    if suf == ".docx":
        return _read_docx(path)
    if suf == ".pdf":
        return _read_pdf(path)
    return None


def load_doc_paths(paths: list[Path | str]) -> tuple[list[str], list[dict]]:
    """显式文件列表 → (documents, doc_metadata)。"""
    documents: list[str] = []
    metadatas: list[dict] = []

    for raw in paths:
        path = Path(raw)
        if not path.is_file():
            logger.warning("不是文件，已跳过: %s", path)
            continue

        suf = path.suffix.lower()
        if suf not in _SUPPORTED:
            if suf in _TABLE_SUFFIXES:
                logger.info(
                    "表格类 %s 应走 sqlite_tool.ingest，RAG 已跳过: %s",
                    path.suffix,
                    path.name,
                )
            else:
                logger.warning("暂不支持类型 %s，已跳过: %s", path.suffix, path)
            continue

        text = _read_by_suffix(path, suf)
        if text is None:
            continue

        resolved = path.resolve()
        documents.append(text)
        metadatas.append(
            {
                "doc_id": _stable_doc_id(resolved),
                "source": resolved.name,
                "source_path": str(resolved),
                "file_type": _file_type(suf),
            }
        )

    logger.info("摄取完成: %s 篇有效文档", len(documents))
    return documents, metadatas


def load_docs_dir(directory: Path | str | None = None) -> tuple[list[str], list[dict]]:
    """扫描目录下 RAG 白名单后缀（不递归）。默认 Config.RAW_DOCS_DIR。"""
    root = Path(directory) if directory is not None else Config.RAW_DOCS_DIR
    if not root.is_dir():
        logger.warning("摄取目录不存在: %s", root)
        return [], []

    paths: list[Path] = []
    for suf in sorted(_SUPPORTED):
        paths.extend(root.glob(f"*{suf}"))
    paths = sorted(
        {p.resolve(): p for p in paths}.values(),
        key=lambda p: p.name.lower(),
    )

    if not paths:
        logger.warning(
            "目录内无 RAG 文档 %s: %s",
            sorted(_SUPPORTED),
            root,
        )
        return [], []

    return load_doc_paths(paths)
