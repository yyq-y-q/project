"""临时文件：上传、限定范围检索、到期清理。"""
from __future__ import annotations

import hashlib
import io
import os
import re
import uuid
from datetime import timedelta
from pathlib import Path
from typing import Any, BinaryIO

from app.auth import Principal, require_permission
from app.errors import TemporaryFileError, ValidationError
from app.governance import (
    audit_event,
    get_temp_file,
    iso_now,
    purge_expired_temp_files,
    register_temp_file,
    utc_now,
)
from app.logging_setup import get_logger

logger = get_logger(__name__)
_ROOT = Path(__file__).resolve().parents[2]
_MAX_BYTES = 8 * 1024 * 1024
_MAX_RESULTS = 30
_ALLOWED = {".txt", ".md", ".markdown", ".csv", ".json", ".log", ".py", ".js", ".ts", ".yaml", ".yml", ".sql", ".docx", ".pdf"}
_SAFE_NAME = re.compile(r"[^a-zA-Z0-9._\-\u4e00-\u9fff]+")


def _root() -> Path:
    path = Path(os.getenv("KB_TEMP_FILE_DIR", str(_ROOT / "data" / "temp_files"))).resolve()
    path.mkdir(parents=True, exist_ok=True)
    return path


def _ttl_minutes() -> int:
    try:
        return min(24 * 60, max(5, int(os.getenv("KB_TEMP_FILE_TTL_MINUTES", "60"))))
    except ValueError:
        return 60


def _safe_name(filename: str) -> str:
    name = Path(filename or "upload").name
    clean = _SAFE_NAME.sub("_", name).strip("._") or "upload"
    return clean[:180]


def _extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix in {".docx"}:
            from docx import Document

            doc = Document(str(path))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(parts)
        if suffix == ".pdf":
            from pypdf import PdfReader

            return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
        raw = path.read_bytes()
        try:
            return raw.decode("utf-8")
        except UnicodeDecodeError:
            return raw.decode("gb18030")
    except Exception as exc:
        raise TemporaryFileError("临时文件无法解析", stage="temp_file.extract", cause=exc) from exc


def cleanup_expired(*, request_id: str, source_ip: str) -> int:
    expired = purge_expired_temp_files()
    count = 0
    for item in expired:
        try:
            Path(item["stored_path"]).unlink(missing_ok=True)
        except OSError as exc:
            logger.warning("expired temp file remove failed path=%s error=%s", item["stored_path"], exc)
        audit_event(
            request_id=request_id,
            action="temp_file.expire",
            status="success",
            operator="system",
            key_id="system",
            source_ip=source_ip,
            target_id=item["file_id"],
            before_hash=item["sha256"],
            failure_reason=None,
        )
        count += 1
    return count


def upload_temp_file(
    *,
    filename: str,
    content_type: str | None,
    stream: BinaryIO,
    principal: Principal,
    request_id: str,
    source_ip: str,
) -> dict[str, Any]:
    require_permission(principal, "temp_file.upload")
    cleanup_expired(request_id=request_id, source_ip=source_ip)
    suffix = Path(filename or "").suffix.lower()
    if suffix not in _ALLOWED:
        raise ValidationError("临时文件类型不受支持", stage="temp_file.upload", details={"suffix": suffix})
    raw = stream.read(_MAX_BYTES + 1)
    if len(raw) > _MAX_BYTES:
        raise ValidationError("临时文件不能超过 8MB", stage="temp_file.upload")
    if not raw:
        raise ValidationError("临时文件不能为空", stage="temp_file.upload")
    file_id = f"tmp_{uuid.uuid4().hex}"
    dest = _root() / f"{file_id}_{_safe_name(filename)}"
    dest.write_bytes(raw)
    created = utc_now()
    expires = created + timedelta(minutes=_ttl_minutes())
    digest = hashlib.sha256(raw).hexdigest()
    register_temp_file({
        "file_id": file_id,
        "stored_path": str(dest),
        "original_name": _safe_name(filename),
        "content_type": content_type,
        "size_bytes": len(raw),
        "sha256": digest,
        "owner_key_id": principal.key_id,
        "owner_operator": principal.operator,
        "created_at": created.isoformat(timespec="seconds"),
        "expires_at": expires.isoformat(timespec="seconds"),
    })
    event_id = audit_event(
        request_id=request_id,
        action="temp_file.upload",
        status="success",
        operator=principal.operator,
        key_id=principal.key_id,
        source_ip=source_ip,
        target_id=file_id,
        after_hash=digest,
    )
    return {
        "ok": True,
        "file_id": file_id,
        "filename": _safe_name(filename),
        "size_bytes": len(raw),
        "sha256": digest,
        "created_at": created.isoformat(timespec="seconds"),
        "expires_at": expires.isoformat(timespec="seconds"),
        "audit_event_id": event_id,
    }


def search_temp_file(
    *,
    file_id: str,
    query: str,
    principal: Principal,
    request_id: str,
    source_ip: str,
) -> dict[str, Any]:
    require_permission(principal, "temp_file.search")
    cleanup_expired(request_id=request_id, source_ip=source_ip)
    text_query = (query or "").strip()
    if not text_query:
        raise ValidationError("query 不能为空", stage="temp_file.search")
    item = get_temp_file(file_id)
    if not item:
        raise TemporaryFileError("临时文件不存在或已过期", stage="temp_file.search")
    if item["owner_key_id"] != principal.key_id and principal.role.value != "admin":
        raise TemporaryFileError("临时文件不属于当前操作者", stage="temp_file.search")
    if item["expires_at"] <= iso_now():
        raise TemporaryFileError("临时文件已过期", stage="temp_file.search")
    text = _extract_text(Path(item["stored_path"]))
    needle = text_query.casefold()
    lines = text.splitlines() or [text]
    hits: list[dict[str, Any]] = []
    for index, line in enumerate(lines, start=1):
        if needle in line.casefold():
            hits.append({"line": index, "text": line[:2000]})
            if len(hits) >= _MAX_RESULTS:
                break
    event_id = audit_event(
        request_id=request_id,
        action="temp_file.search",
        status="success",
        operator=principal.operator,
        key_id=principal.key_id,
        source_ip=source_ip,
        target_id=file_id,
        before_hash=item["sha256"],
    )
    return {
        "ok": True,
        "file_id": file_id,
        "query": text_query,
        "matches": hits,
        "match_count": len(hits),
        "truncated": len(hits) == _MAX_RESULTS,
        "expires_at": item["expires_at"],
        "audit_event_id": event_id,
    }
